import io
import json
import os
import re
from typing import Dict, List, Tuple, Optional

import pandas as pd
import requests
import streamlit as st
from bs4 import BeautifulSoup

try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except Exception:
    pdf_extract_text = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    from openai import AzureOpenAI
except Exception:
    AzureOpenAI = None


# ============================================================
# Settings / Constants
# ============================================================

MAX_FILE_BYTES = 25 * 1024 * 1024
LLM_SAMPLE_CHARS = 5000

# Sanity caps: prevent absurd values when extraction goes sideways
MAX_REASONABLE_WORDS_PER_DOC = 200_000
MAX_REASONABLE_MINUTES_PER_DOC = 1_000

# Heuristics for non-text
MIN_PER_PDF_PAGE = 3.5
MIN_PER_PPT_SLIDE = 2.0

# Prefer formats (within a single item) when same filename appears as PDF + PPTX etc.
PREFERRED_EXT_ORDER = ["pdf", "docx", "doc", "txt", "html", "htm", "pptx", "ppt"]


def get_secret(name: str, default=None):
    try:
        return st.secrets[name]
    except Exception:
        return os.getenv(name, default)


CANVAS_BASE = (get_secret("CANVAS_BASE_URL", "") or "").rstrip("/")
CANVAS_TOKEN = get_secret("CANVAS_API_TOKEN", "") or ""

AZ_ENDPOINT = get_secret("AZURE_OPENAI_ENDPOINT", "") or ""
AZ_API_KEY = get_secret("AZURE_OPENAI_API_KEY", "") or ""
AZ_MODEL = get_secret("AZURE_OPENAI_MODEL", "") or ""
AZ_API_VERSION = get_secret("AZURE_OPENAI_API_VERSION", "2024-02-15-preview") or "2024-02-15-preview"


# ============================================================
# Utility: formatting
# ============================================================

def minutes_to_hhmm(minutes: float) -> str:
    if minutes is None:
        return "00:00"
    try:
        total = int(round(float(minutes)))
    except Exception:
        return "00:00"
    h, m = divmod(max(total, 0), 60)
    return f"{h:02d}:{m:02d}"


def hhmmss_to_seconds(hhmmss: str) -> int:
    parts = (hhmmss or "").strip().split(":")
    if len(parts) != 3:
        return 0
    try:
        h, m, s = [int(x) for x in parts]
    except Exception:
        return 0
    return max(0, h * 3600 + m * 60 + s)


# ============================================================
# Utility: filename / type
# ============================================================

def file_ext(name: str) -> str:
    name = (name or "").strip().lower()
    m = re.search(r"\.([a-z0-9]+)$", name)
    return m.group(1) if m else ""


def file_stem(name: str) -> str:
    name = (name or "").strip()
    # remove extension
    name = re.sub(r"\.[A-Za-z0-9]+$", "", name)
    return name.strip().lower()


def rank_ext(ext: str) -> int:
    try:
        return PREFERRED_EXT_ORDER.index(ext)
    except ValueError:
        return 999


def dedupe_linked_files_by_stem(metas: List[dict]) -> List[dict]:
    """
    Dedupe within one item: if same doc exists as PDF + PPTX, prefer PDF.
    Keyed by filename stem.
    """
    best: Dict[str, dict] = {}
    for meta in metas:
        name = meta.get("display_name") or meta.get("filename") or ""
        stem = file_stem(name) or f"id:{meta.get('id')}"
        ext = file_ext(name)

        if stem not in best:
            best[stem] = meta
            continue

        cur = best[stem]
        cur_name = cur.get("display_name") or cur.get("filename") or ""
        cur_ext = file_ext(cur_name)

        if rank_ext(ext) < rank_ext(cur_ext):
            best[stem] = meta

    return list(best.values())


def is_text_like_content_type(ct: str) -> bool:
    ct = (ct or "").lower()
    return ct.startswith("text/") or any(x in ct for x in ["json", "xml", "html"])


# ============================================================
# Canvas HTTP helpers
# ============================================================

def canvas_headers():
    if not CANVAS_TOKEN:
        raise RuntimeError("Missing CANVAS_API_TOKEN.")
    return {"Authorization": f"Bearer {CANVAS_TOKEN}"}


def canvas_get(url: str, params=None) -> List[dict]:
    """Pagination-aware GET for Canvas list endpoints."""
    out = []
    while url:
        r = requests.get(url, headers=canvas_headers(), params=params, timeout=30)
        r.raise_for_status()
        data = r.json()
        if isinstance(data, list):
            out.extend(data)
        else:
            out.append(data)

        link = r.headers.get("Link", "")
        next_url = None
        for part in link.split(","):
            if 'rel="next"' in part:
                m = re.search(r"<([^>]+)>", part)
                if m:
                    next_url = m.group(1)
        url = next_url
        params = None
    return out


@st.cache_data(show_spinner=False)
def fetch_url_bytes(url: str, max_bytes: int) -> Tuple[bytes, str]:
    """
    Fetch bytes from a Canvas signed URL. Returns (data, detected_content_type).
    """
    r = requests.get(
        url,
        headers=canvas_headers(),
        timeout=60,
        allow_redirects=True,
        stream=True,
    )
    r.raise_for_status()
    ct = (r.headers.get("Content-Type") or "").split(";")[0].strip().lower()
    data = r.content[:max_bytes]
    return data, ct


# ============================================================
# Canvas API wrappers
# ============================================================

def get_modules_with_items(course_id: int) -> List[dict]:
    url = f"{CANVAS_BASE}/api/v1/courses/{course_id}/modules"
    mods = canvas_get(url, params={"include[]": "items", "per_page": 100})
    items = []
    for mod in mods:
        for it in mod.get("items", []):
            items.append(
                {
                    "module_name": mod.get("name", ""),
                    "module_position": mod.get("position", 0),
                    "item_type": it.get("type", ""),
                    "title": it.get("title", ""),
                    "html_url": it.get("html_url", ""),
                    "content_id": it.get("content_id"),
                    "page_url": it.get("page_url"),
                    "content_details": it.get("content_details", {}),
                    "item_key": f"{it.get('type','')}::{it.get('id')}",
                }
            )
    return items


def get_page_body(course_id: int, page_url: str) -> str:
    url = f"{CANVAS_BASE}/api/v1/courses/{course_id}/pages/{page_url}"
    r = requests.get(url, headers=canvas_headers(), timeout=30)
    r.raise_for_status()
    return r.json().get("body", "") or ""


def get_assignment(course_id: int, assignment_id: int) -> dict:
    url = f"{CANVAS_BASE}/api/v1/courses/{course_id}/assignments/{assignment_id}"
    r = requests.get(url, headers=canvas_headers(), timeout=30)
    r.raise_for_status()
    return r.json()


def get_discussion(course_id: int, topic_id: int) -> dict:
    url = f"{CANVAS_BASE}/api/v1/courses/{course_id}/discussion_topics/{topic_id}"
    r = requests.get(url, headers=canvas_headers(), timeout=30)
    r.raise_for_status()
    return r.json()


def get_quiz(course_id: int, quiz_id: int) -> dict:
    url = f"{CANVAS_BASE}/api/v1/courses/{course_id}/quizzes/{quiz_id}"
    r = requests.get(url, headers=canvas_headers(), timeout=30)
    r.raise_for_status()
    return r.json()


# ---------------- New Quizzes (default) ----------------


def get_new_quiz(course_id: int, assignment_id: int) -> dict:
    """Fetch a single New Quiz.

    Canvas New Quizzes API uses the *assignment_id* as the quiz identifier.
    Endpoint: GET /api/quiz/v1/courses/:course_id/quizzes/:assignment_id
    """
    url = f"{CANVAS_BASE}/api/quiz/v1/courses/{course_id}/quizzes/{assignment_id}"
    r = requests.get(url, headers=canvas_headers(), timeout=30)
    r.raise_for_status()
    return r.json()


def list_new_quiz_items(course_id: int, assignment_id: int) -> List[dict]:
    """List items in a New Quiz.

    Endpoint: GET /api/quiz/v1/courses/:course_id/quizzes/:assignment_id/items
    """
    url = f"{CANVAS_BASE}/api/quiz/v1/courses/{course_id}/quizzes/{assignment_id}/items"
    return canvas_get(url, params={"per_page": 100})


def get_course_file(course_id: int, file_id: int) -> dict:
    url = f"{CANVAS_BASE}/api/v1/courses/{course_id}/files/{file_id}"
    r = requests.get(url, headers=canvas_headers(), timeout=30)
    r.raise_for_status()
    return r.json()


# ============================================================
# HTML/Text extraction
# ============================================================

def strip_html_to_text(html: str) -> str:
    soup = BeautifulSoup(html or "", "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator=" ")
    return re.sub(r"\s+", " ", text).strip()


def words_from_text(text: str) -> int:
    """
    Ignore 1-character tokens to prevent PDF spaced-letter artifacts
    from inflating "word" counts.
    """
    if not text:
        return 0
    return len(re.findall(r"\b[\w']{2,}\b", text))


def detect_videos_from_html(html: str) -> List[dict]:
    videos = []
    if not html:
        return videos
    soup = BeautifulSoup(html, "html.parser")

    for tag in soup.find_all(["iframe", "video", "embed"]):
        src = tag.get("src") or tag.get("data-src") or ""
        if not src:
            continue
        title = tag.get("title") or tag.get("aria-label") or "Embedded Video"
        videos.append({"src": src, "title": title})

    for a in soup.find_all("a", href=True):
        href = a["href"]
        if any(dom in href for dom in ["youtube.com", "youtu.be", "vimeo.com", "echo360", "panopto", "kaltura"]):
            title = a.get_text(strip=True) or "Linked Video"
            videos.append({"src": href, "title": title})

    return videos


def detect_canvas_file_ids_from_html(html: str) -> List[int]:
    """
    Canvas-only: extract /files/<id> from links.
    Ignores any URLs that aren't a Canvas file link.
    """
    if not html:
        return []
    soup = BeautifulSoup(html, "html.parser")
    ids = set()
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if "/files/" not in href:
            continue
        m = re.search(r"/files/(\d+)", href)
        if m:
            ids.add(int(m.group(1)))
    return sorted(ids)


# ============================================================
# File extraction (local) — Canvas hosted only
# ============================================================

def extract_text_from_canvas_file(
    file_url: str,
    filename: str,
    content_type_hint: str,
    max_bytes: int,
) -> Tuple[str, int, str]:
    """
    Returns (text, pages_or_slides, detected_ct).
    PPTX: returns empty text + slide_count (never word-count PPTX).
    Unknown binary: returns empty text.
    """
    if not file_url:
        return "", 0, ""

    data, detected_ct = fetch_url_bytes(file_url, max_bytes)
    ct = (content_type_hint or detected_ct or "").split(";")[0].strip().lower()
    ext = file_ext(filename)

    # PPTX: never extract words; slide heuristic only
    if ext in ("pptx", "ppt") or "powerpoint" in ct:
        if Presentation is not None:
            try:
                prs = Presentation(io.BytesIO(data))
                return "", len(prs.slides), ct or "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            except Exception:
                return "", 0, ct
        return "", 0, ct

    # PDF
    if ext == "pdf" or ("pdf" in ct):
        if pdf_extract_text is None:
            return "", 0, ct
        try:
            text = pdf_extract_text(io.BytesIO(data))
            pages = text.count("\f") or 0
            return text, pages, ct or "application/pdf"
        except Exception:
            return "", 0, ct

    # DOCX
    if ext in ("docx", "doc") or ("word" in ct) or ("docx" in ct):
        if Document is None:
            return "", 0, ct
        try:
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            return text, 0, ct or "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception:
            return "", 0, ct

    # Text-like
    if is_text_like_content_type(ct) or ext in ("txt", "html", "htm", "csv", "md"):
        try:
            return data.decode("utf-8", errors="ignore"), 0, ct
        except Exception:
            return "", 0, ct

    # Unknown/binary => no text
    return "", 0, ct


def minutes_fallback_for_nontext(pages_or_slides: int, filename: str, content_type: str) -> float:
    ext = file_ext(filename)
    ct = (content_type or "").lower()
    if pages_or_slides <= 0:
        # conservative default if we can't measure
        return 10.0
    if ext in ("pptx", "ppt") or ("powerpoint" in ct):
        return float(pages_or_slides) * MIN_PER_PPT_SLIDE
    return float(pages_or_slides) * MIN_PER_PDF_PAGE


# ============================================================
# LLM difficulty — SAMPLE ONLY
# ============================================================

def azure_llm_client():
    if AzureOpenAI is None:
        raise RuntimeError("openai SDK not installed.")
    return AzureOpenAI(api_key=AZ_API_KEY, azure_endpoint=AZ_ENDPOINT.rstrip("/"), api_version=AZ_API_VERSION)


def default_difficulty() -> Dict:
    return {"label": "average", "wpm_factor": 1.0, "notes": "default (LLM off/failed)"}


def azure_llm_difficulty_sample(sample_text: str) -> Dict:
    """
    Uses ONLY a sample (<=5000 chars). Returns {label, wpm_factor, notes}.
    """
    if not (AZ_ENDPOINT and AZ_API_KEY and AZ_MODEL):
        return default_difficulty()

    client = azure_llm_client()
    sys_msg = (
        "You are a reading difficulty estimator for college coursework. "
        "Return ONLY JSON with keys:\n"
        "label: one of ['very_easy','easy','average','hard','very_hard']\n"
        "wpm_factor: float multiplier relative to base reading speed\n"
        "notes: short explanation\n"
        "Use: very_easy=1.3, easy=1.15, average=1.0, hard=0.8, very_hard=0.65."
    )
    user_msg = f"Estimate difficulty for this sample:\n\n{sample_text}"

    try:
        cc = client.chat.completions.create(
            model=AZ_MODEL,
            messages=[{"role": "system", "content": sys_msg}, {"role": "user", "content": user_msg}],
            temperature=0,
            response_format={"type": "json_object"},
        )
        data = json.loads(cc.choices[0].message.content)
        return {
            "label": data.get("label", "average"),
            "wpm_factor": float(data.get("wpm_factor", 1.0)),
            "notes": data.get("notes", ""),
        }
    except Exception:
        # fallback parse
        try:
            cc = client.chat.completions.create(
                model=AZ_MODEL,
                messages=[{"role": "system", "content": sys_msg}, {"role": "user", "content": user_msg}],
                temperature=0,
            )
            raw = cc.choices[0].message.content or ""
            m = re.search(r"{.*}", raw, flags=re.DOTALL)
            if m:
                data = json.loads(m.group(0))
                return {
                    "label": data.get("label", "average"),
                    "wpm_factor": float(data.get("wpm_factor", 1.0)),
                    "notes": data.get("notes", "fallback parse"),
                }
        except Exception:
            pass
        return default_difficulty()


def reading_minutes_from_words(words: int, base_wpm: int, wpm_factor: float) -> float:
    wpm = max(80.0, float(base_wpm) * max(0.1, float(wpm_factor)))
    return float(words) / wpm


# ============================================================
# DO-time estimation (bounded + explainable)
# ============================================================


def clamp(v: float, lo: float, hi: float) -> float:
    try:
        v = float(v)
    except Exception:
        v = 0.0
    return max(lo, min(hi, v))


def default_task_complexity() -> Dict[str, str]:
    return {"task_complexity": "standard", "notes": "default (LLM off/failed)"}


def azure_llm_task_complexity(text: str, item_type: str, level: str) -> Dict[str, str]:
    """Classify task complexity (NOT minutes).

    Returns {task_complexity: one of simple|standard|complex|very_complex, notes: str}
    """
    if not (AZ_ENDPOINT and AZ_API_KEY and AZ_MODEL) or AzureOpenAI is None:
        return default_task_complexity()

    client = azure_llm_client()
    sys_msg = (
        "You are classifying the complexity of a single Canvas course task. "
        "Return ONLY JSON with keys: task_complexity, notes. "
        "task_complexity must be one of: simple, standard, complex, very_complex. "
        "Do NOT estimate minutes."
    )
    user_msg = (
        f"Item type: {item_type}\nStudent level: {level}\n\n"
        "Classify complexity of completing the task (excluding reading/watching).\n\n"
        f"{(text or '')[:8000]}"
    )

    try:
        cc = client.chat.completions.create(
            model=AZ_MODEL,
            messages=[{"role": "system", "content": sys_msg}, {"role": "user", "content": user_msg}],
            temperature=0,
            response_format={"type": "json_object"},
        )
        data = json.loads(cc.choices[0].message.content)
        tc = (data.get("task_complexity") or "standard").strip().lower()
        if tc not in TASK_COMPLEXITY_MULT:
            tc = "standard"
        return {"task_complexity": tc, "notes": (data.get("notes") or "").strip()}
    except Exception:
        return default_task_complexity()


def complexity_multiplier(task_complexity: str) -> float:
    tc = (task_complexity or "standard").strip().lower()
    return float(TASK_COMPLEXITY_MULT.get(tc, 1.0))


def assignment_base_minutes(prompt_words: int, level: str) -> float:
    """Explainable baseline based on instruction length bands."""
    lvl_factor = 1.0 if level.lower().startswith("under") else 1.25
    if prompt_words < 150:
        base = 30.0
    elif prompt_words < 600:
        base = 60.0
    else:
        base = 120.0
    return base * lvl_factor


def discussion_component_minutes(prompt_words: int, level: str) -> Dict[str, float]:
    """Component model for discussions based on prompt length."""
    lvl_factor = 1.0 if level.lower().startswith("under") else 1.25

    if prompt_words < 150:
        initial = 18.0
        read_peers = 10.0
        replies = 7.0
    elif prompt_words < 400:
        initial = 25.0
        read_peers = 12.0
        replies = 10.0
    else:
        initial = 35.0
        read_peers = 15.0
        replies = 12.0

    return {
        "initial_post": initial * lvl_factor,
        "read_peers": read_peers * lvl_factor,
        "replies": replies * lvl_factor,
    }


def new_quiz_time_limit_minutes(new_quiz: dict) -> Optional[float]:
    """If New Quiz has an explicit time limit, return it in minutes."""
    qs = (new_quiz or {}).get("quiz_settings") or {}
    if qs.get("has_time_limit") is True:
        secs = qs.get("session_time_limit_in_seconds")
        try:
            secs = int(secs)
        except Exception:
            secs = None
        if secs and secs > 0:
            return float(secs) / 60.0
    return None


def classic_quiz_time_limit_minutes(classic_quiz: dict) -> Optional[float]:
    t = (classic_quiz or {}).get("time_limit")
    try:
        t = float(t)
    except Exception:
        t = None
    return t if (t and t > 0) else None


def estimate_new_quiz_minutes_from_items(items: List[dict], fallback_question_count: Optional[int] = None) -> float:
    """Sum per-question seconds using New Quiz Items interaction_type_slug."""
    total_seconds = 0
    counted = 0

    # New Quizzes interaction_type_slug values are documented in New Quiz Items API
    # (e.g., 'choice', 'true-false', 'matching', 'essay', 'numeric', etc.).
    seconds_by_slug = {
        "choice": 75,
        "true-false": 60,
        "multi-answer": 90,
        "matching": 105,
        "categorization": 120,
        "ordering": 120,
        "rich-fill-blank": 150,
        "numeric": 150,
        "formula": 180,
        "file-upload": 300,
        "hot-spot": 150,
        "essay": 900,
    }

    for it in items or []:
        entry = (it or {}).get("entry") or {}
        slug = (entry.get("interaction_type_slug") or "").strip().lower()
        if not slug:
            continue
        counted += 1
        total_seconds += int(seconds_by_slug.get(slug, QUIZ_SECONDS_DEFAULT))

    if counted > 0:
        return max(5.0, total_seconds / 60.0)

    # Fallback if items can't be read (permissions / quiz type / etc.)
    qcount = fallback_question_count or 0
    try:
        qcount = int(qcount)
    except Exception:
        qcount = 0
    if qcount > 0:
        return max(5.0, (qcount * QUIZ_SECONDS_DEFAULT) / 60.0)
    return 10.0


# ============================================================
# App
# ============================================================

def main():
    st.set_page_config(page_title="Course Load Estimator", layout="wide")
    st.title("📚 Course Load Estimator")

    # session state
    st.session_state.setdefault("items", [])
    st.session_state.setdefault("results", [])
    st.session_state.setdefault("pending_videos", {})

    # Sidebar config
    st.sidebar.header("Configuration")
    course_id = st.sidebar.text_input("Canvas Course ID", value="")
    level = st.sidebar.selectbox("Student Level", ["Undergraduate", "Graduate"])
    base_wpm = st.sidebar.slider("Base Reading Speed (words per minute)", 150, 350, 200, 10)
    use_llm = st.sidebar.checkbox("Use Azure OpenAI for difficulty & DO time", value=True)
    debug_breakdown = st.sidebar.checkbox("Debug read-time breakdown", value=False)
    debug_do_breakdown = st.sidebar.checkbox("Debug DO-time breakdown", value=False)

    # Status
    st.sidebar.markdown("### Canvas status")
    if not (CANVAS_BASE and CANVAS_TOKEN):
        st.sidebar.error("Canvas secrets missing or incomplete.")
    else:
        st.sidebar.success("Canvas configured.")

    st.sidebar.markdown("### Azure OpenAI status")
    if not (AZ_ENDPOINT and AZ_API_KEY and AZ_MODEL):
        st.sidebar.warning("Azure OpenAI secrets missing or incomplete.")
    else:
        st.sidebar.success("Azure OpenAI configured.")

    # KPIs (hh:mm) from current results
    if st.session_state.get("results"):
        df_all = pd.DataFrame(st.session_state["results"])
        total_read = df_all.get("read_min", pd.Series(dtype=float)).sum()
        total_watch = df_all.get("watch_min", pd.Series(dtype=float)).sum()
        total_do = df_all.get("do_min", pd.Series(dtype=float)).sum()
        total_total = df_all.get("total_min", pd.Series(dtype=float)).sum()

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total Read (hh:mm)", minutes_to_hhmm(total_read))
        c2.metric("Total Watch (hh:mm)", minutes_to_hhmm(total_watch))
        c3.metric("Total Do (hh:mm)", minutes_to_hhmm(total_do))
        c4.metric("Total Workload (hh:mm)", minutes_to_hhmm(total_total))

    st.markdown(
        """
This estimator calculates:
- **READ**: Canvas page text + Canvas-hosted linked documents (PDF/DOCX/etc).  
  - Word counts extracted locally.
  - Difficulty factor from Azure OpenAI using **only the first 5,000 characters**.
- **WATCH**: videos detected in content; you enter duration **per video**.
- **DO**: assignments/discussions/quizzes (LLM or heuristic).
"""
    )

    # 1) Scan
    st.header("1) Scan Course")
    if st.button("Scan course modules & items", type="primary"):
        if not course_id:
            st.error("Enter a Canvas Course ID.")
        elif not (CANVAS_BASE and CANVAS_TOKEN):
            st.error("Canvas configuration not set.")
        else:
            try:
                with st.spinner("Fetching modules and items from Canvas..."):
                    items = get_modules_with_items(int(course_id))
                st.session_state["items"] = items
                st.session_state["results"] = []
                st.session_state["pending_videos"] = {}
                st.success(f"Found {len(items)} module items.")
            except Exception as e:
                st.error(f"Canvas API error: {e}")

    if st.session_state["items"]:
        st.write(f"Total items discovered: **{len(st.session_state['items'])}**")
        with st.expander("Preview raw module items"):
            st.json(st.session_state["items"])

    # 2) Process
    st.header("2) Estimate READ and DO time")

    if st.button("Process items for workload"):
        items = st.session_state.get("items", [])
        if not items:
            st.warning("No items scanned yet.")
        else:
            if use_llm and not (AZ_ENDPOINT and AZ_API_KEY and AZ_MODEL):
                st.error("Azure OpenAI not configured.")
                return

            results = []
            debug_rows: List[dict] = []
            debug_do_rows: List[dict] = []

            for it in items:
                item_type = it.get("item_type", "")
                title = it.get("title", "")
                html_url = it.get("html_url", "")
                item_key = it.get("item_key", "")

                read_min = 0.0
                watch_min = 0.0
                do_min = 0.0
                task_complexity = "standard"
                do_notes = ""

                # -------- Pages / Assignments / Discussions --------
                if item_type in ("Page", "Assignment", "Discussion"):
                    try:
                        if item_type == "Page":
                            body_html = get_page_body(int(course_id), it.get("page_url"))
                        elif item_type == "Assignment":
                            a = get_assignment(int(course_id), it.get("content_id"))
                            body_html = a.get("description", "") or ""
                        else:
                            d = get_discussion(int(course_id), it.get("content_id"))
                            body_html = d.get("message", "") or ""
                    except Exception:
                        body_html = ""

                    # detect videos (per-video entry later)
                    vids = detect_videos_from_html(body_html)
                    for idx, v in enumerate(vids, start=1):
                        v_key = f"{item_key}::embed::{idx}"
                        st.session_state["pending_videos"].setdefault(
                            v_key,
                            {
                                "title": v.get("title", "Video"),
                                "src": v.get("src", ""),
                                "hhmmss": "00:00:00",
                                "seconds": 0,
                                "item_key": item_key,
                            },
                        )

                    # Page text -> local words
                    page_text = strip_html_to_text(body_html)
                    page_words = words_from_text(page_text)

                    if page_words > 0:
                        if use_llm:
                            diff = azure_llm_difficulty_sample(page_text[:LLM_SAMPLE_CHARS])
                            wpm_factor = diff.get("wpm_factor", 1.0)
                        else:
                            diff = default_difficulty()
                            wpm_factor = 1.0

                        page_minutes = reading_minutes_from_words(page_words, base_wpm, wpm_factor)
                        read_min += page_minutes

                        if debug_breakdown:
                            debug_rows.append({
                                "item": title,
                                "component": "page_text",
                                "name": "(page text)",
                                "filename": "",
                                "content_type": "text/html",
                                "words": page_words,
                                "minutes": page_minutes,
                                "difficulty_label": diff.get("label"),
                                "wpm_factor": wpm_factor,
                                "note": "",
                            })

                    # Linked Canvas files (Canvas-only)
                    file_ids = detect_canvas_file_ids_from_html(body_html)
                    metas: List[dict] = []
                    for fid in file_ids:
                        try:
                            metas.append(get_course_file(int(course_id), fid))
                        except Exception:
                            continue

                    # Dedupe within this item; prefer PDF
                    metas = dedupe_linked_files_by_stem(metas)

                    for meta in metas:
                        filename = meta.get("display_name") or meta.get("filename") or ""
                        ct_hint = (meta.get("content-type") or meta.get("content_type") or "").lower()
                        file_url = meta.get("url") or meta.get("download_url")

                        if not file_url:
                            continue

                        text, pages_or_slides, detected_ct = extract_text_from_canvas_file(
                            file_url=file_url,
                            filename=filename,
                            content_type_hint=ct_hint,
                            max_bytes=MAX_FILE_BYTES,
                        )
                        ct = detected_ct or ct_hint
                        ext = file_ext(filename)

                        # PPTX: always heuristic (no words)
                        if ext in ("pptx", "ppt") or "powerpoint" in (ct or ""):
                            doc_minutes = minutes_fallback_for_nontext(pages_or_slides, filename, ct)
                            read_min += doc_minutes
                            if debug_breakdown:
                                debug_rows.append({
                                    "item": title,
                                    "component": "linked_doc",
                                    "name": meta.get("display_name") or meta.get("filename") or f"file:{meta.get('id')}",
                                    "filename": filename,
                                    "content_type": ct,
                                    "words": 0,
                                    "minutes": doc_minutes,
                                    "difficulty_label": None,
                                    "wpm_factor": None,
                                    "note": f"pptx/slides heuristic ({pages_or_slides} slides)",
                                })
                            continue

                        # Text-bearing doc -> local words
                        doc_words = words_from_text(text)

                        # Sanity: if extraction went nuts or doc is empty, fallback
                        if doc_words <= 0:
                            doc_minutes = minutes_fallback_for_nontext(pages_or_slides, filename, ct)
                            read_min += doc_minutes
                            if debug_breakdown:
                                debug_rows.append({
                                    "item": title,
                                    "component": "linked_doc",
                                    "name": meta.get("display_name") or meta.get("filename") or f"file:{meta.get('id')}",
                                    "filename": filename,
                                    "content_type": ct,
                                    "words": doc_words,
                                    "minutes": doc_minutes,
                                    "difficulty_label": None,
                                    "wpm_factor": None,
                                    "note": "no-text fallback",
                                })
                            continue

                        if doc_words > MAX_REASONABLE_WORDS_PER_DOC:
                            doc_minutes = minutes_fallback_for_nontext(pages_or_slides, filename, ct)
                            read_min += doc_minutes
                            if debug_breakdown:
                                debug_rows.append({
                                    "item": title,
                                    "component": "linked_doc",
                                    "name": meta.get("display_name") or meta.get("filename") or f"file:{meta.get('id')}",
                                    "filename": filename,
                                    "content_type": ct,
                                    "words": doc_words,
                                    "minutes": doc_minutes,
                                    "difficulty_label": None,
                                    "wpm_factor": None,
                                    "note": f"sanity fallback (words>{MAX_REASONABLE_WORDS_PER_DOC})",
                                })
                            continue

                        # Difficulty from SAMPLE ONLY (first 5000 chars)
                        if use_llm:
                            diff = azure_llm_difficulty_sample((text or "")[:LLM_SAMPLE_CHARS])
                            wpm_factor = diff.get("wpm_factor", 1.0)
                        else:
                            diff = default_difficulty()
                            wpm_factor = 1.0

                        doc_minutes = reading_minutes_from_words(doc_words, base_wpm, wpm_factor)

                        if doc_minutes > MAX_REASONABLE_MINUTES_PER_DOC:
                            doc_minutes = minutes_fallback_for_nontext(pages_or_slides, filename, ct)
                            note = f"sanity fallback (minutes>{MAX_REASONABLE_MINUTES_PER_DOC})"
                        else:
                            note = ""

                        read_min += doc_minutes

                        if debug_breakdown:
                            debug_rows.append({
                                "item": title,
                                "component": "linked_doc",
                                "name": meta.get("display_name") or meta.get("filename") or f"file:{meta.get('id')}",
                                "filename": filename,
                                "content_type": ct,
                                "words": doc_words,
                                "minutes": doc_minutes,
                                "difficulty_label": diff.get("label"),
                                "wpm_factor": wpm_factor,
                                "note": note,
                            })

                    # DO time for assignments/discussions
                    if item_type in ("Assignment", "Discussion"):
                        comp = azure_llm_task_complexity(page_text, item_type, level) if use_llm else default_task_complexity()
                        task_complexity = comp.get("task_complexity", "standard")
                        do_notes = comp.get("notes", "")
                        mult = complexity_multiplier(task_complexity)

                        if item_type == "Assignment":
                            base_do = assignment_base_minutes(page_words, level)
                            do_min = clamp(base_do * mult, ASSIGNMENT_MIN_MIN, ASSIGNMENT_MAX_MIN)
                            if debug_do_breakdown:
                                debug_do_rows.append({
                                    "item": title,
                                    "type": item_type,
                                    "prompt_words": page_words,
                                    "task_complexity": task_complexity,
                                    "multiplier": mult,
                                    "base_do_min": round(base_do, 2),
                                    "do_min": round(do_min, 2),
                                    "notes": do_notes,
                                })
                        else:
                            comps = discussion_component_minutes(page_words, level)
                            base_do = float(sum(comps.values()))
                            do_min = clamp(base_do * mult, DISCUSSION_MIN_MIN, DISCUSSION_MAX_MIN)
                            if debug_do_breakdown:
                                row = {
                                    "item": title,
                                    "type": item_type,
                                    "prompt_words": page_words,
                                    "task_complexity": task_complexity,
                                    "multiplier": mult,
                                    "base_do_min": round(base_do, 2),
                                    "do_min": round(do_min, 2),
                                    "notes": do_notes,
                                }
                                row.update({k: round(v, 2) for k, v in comps.items()})
                                debug_do_rows.append(row)

                # -------- File module items (Canvas file items) --------
                elif item_type == "File":
                    cd = it.get("content_details") or {}
                    file_url = cd.get("url")
                    filename = cd.get("display_name") or cd.get("filename") or title or ""
                    ct_hint = (cd.get("content_type") or "").lower()

                    if file_url:
                        text, pages_or_slides, detected_ct = extract_text_from_canvas_file(
                            file_url=file_url,
                            filename=filename,
                            content_type_hint=ct_hint,
                            max_bytes=MAX_FILE_BYTES,
                        )
                        ct = detected_ct or ct_hint
                        ext = file_ext(filename)

                        if ext in ("pptx", "ppt") or "powerpoint" in (ct or ""):
                            read_min = minutes_fallback_for_nontext(pages_or_slides, filename, ct)
                        else:
                            w = words_from_text(text)
                            if w <= 0 or w > MAX_REASONABLE_WORDS_PER_DOC:
                                read_min = minutes_fallback_for_nontext(pages_or_slides, filename, ct)
                            else:
                                if use_llm:
                                    diff = azure_llm_difficulty_sample((text or "")[:LLM_SAMPLE_CHARS])
                                    wpm_factor = diff.get("wpm_factor", 1.0)
                                else:
                                    wpm_factor = 1.0
                                read_min = reading_minutes_from_words(w, base_wpm, wpm_factor)
                                if read_min > MAX_REASONABLE_MINUTES_PER_DOC:
                                    read_min = minutes_fallback_for_nontext(pages_or_slides, filename, ct)

                # -------- Quiz --------
                elif item_type == "Quiz":
                    q_meta = it.get("content_details") or {}
                    quiz_id = it.get("content_id")

                    # Default to New Quizzes endpoints first; fall back to Classic.
                    new_quiz = None
                    classic_quiz = None

                    # Complexity classification context (used unless explicit time limit exists)
                    if use_llm:
                        # We'll classify off whatever instructions we can fetch below.
                        pass

                    # Try New Quizzes API (quiz id is assignment_id)
                    try:
                        if quiz_id:
                            new_quiz = get_new_quiz(int(course_id), int(quiz_id))
                    except Exception:
                        new_quiz = None

                    if new_quiz is not None:
                        # If time limit explicitly set, ALWAYS default to it
                        tl = new_quiz_time_limit_minutes(new_quiz)
                        instructions = strip_html_to_text(new_quiz.get("instructions", "") or "")
                        comp = azure_llm_task_complexity(instructions, "Quiz", level) if use_llm else default_task_complexity()
                        task_complexity = comp.get("task_complexity", "standard")
                        do_notes = comp.get("notes", "")
                        mult = complexity_multiplier(task_complexity)

                        if tl is not None:
                            do_min = float(tl)
                            do_notes = (do_notes + " | used explicit time limit").strip(" |")
                        else:
                            try:
                                items = list_new_quiz_items(int(course_id), int(quiz_id))
                            except Exception:
                                items = []
                            base_do = estimate_new_quiz_minutes_from_items(
                                items,
                                fallback_question_count=q_meta.get("question_count"),
                            )
                            do_min = max(5.0, base_do * mult)

                        if debug_do_breakdown:
                            debug_do_rows.append({
                                "item": title,
                                "type": item_type,
                                "prompt_words": words_from_text(instructions),
                                "task_complexity": task_complexity,
                                "multiplier": mult,
                                "base_do_min": round(float(tl if tl is not None else base_do), 2),
                                "do_min": round(do_min, 2),
                                "notes": do_notes,
                                "quiz_api": "new_quizzes",
                            })

                    else:
                        # Classic fallback
                        try:
                            if quiz_id:
                                classic_quiz = get_quiz(int(course_id), int(quiz_id))
                        except Exception:
                            classic_quiz = None

                        desc = strip_html_to_text((classic_quiz or {}).get("description", "") or "")
                        comp = azure_llm_task_complexity(desc, "Quiz", level) if use_llm else default_task_complexity()
                        task_complexity = comp.get("task_complexity", "standard")
                        do_notes = comp.get("notes", "")
                        mult = complexity_multiplier(task_complexity)

                        tl = classic_quiz_time_limit_minutes(classic_quiz or {})
                        if tl is not None:
                            do_min = float(tl)
                            do_notes = (do_notes + " | used explicit time limit").strip(" |")
                        else:
                            qcount = (classic_quiz or {}).get("question_count") or q_meta.get("question_count") or 0
                            try:
                                qcount = int(qcount)
                            except Exception:
                                qcount = 0
                            base_do = max(5.0, (qcount * QUIZ_SECONDS_DEFAULT) / 60.0) if qcount > 0 else 10.0
                            do_min = max(5.0, base_do * mult)

                        if debug_do_breakdown:
                            debug_do_rows.append({
                                "item": title,
                                "type": item_type,
                                "prompt_words": words_from_text(desc),
                                "task_complexity": task_complexity,
                                "multiplier": mult,
                                "base_do_min": round(float(tl if tl is not None else base_do), 2),
                                "do_min": round(do_min, 2),
                                "notes": do_notes,
                                "quiz_api": "classic_fallback",
                            })

                # -------- External link video item --------
                else:
                    # Do not fetch anything external; only detect if it's a video and ask for manual duration
                    if any(dom in (html_url or "") for dom in ("youtube", "youtu.be", "vimeo", "echo360", "panopto", "kaltura")):
                        v_key = f"{item_key}::external"
                        st.session_state["pending_videos"].setdefault(
                            v_key,
                            {
                                "title": title or "External Video",
                                "src": html_url,
                                "hhmmss": "00:00:00",
                                "seconds": 0,
                                "item_key": item_key,
                            },
                        )

                total_min = float(read_min) + float(watch_min) + float(do_min)

                results.append(
                    {
                        "module": it.get("module_name", ""),
                        "module_position": it.get("module_position", 0),
                        "title": title,
                        "type": item_type,
                        "url": html_url,
                        "item_key": item_key,
                        "read_min": round(read_min, 2),
                        "watch_min": round(watch_min, 2),
                        "do_min": round(do_min, 2),
                        "task_complexity": task_complexity,
                        "do_notes": do_notes,
                        "total_min": round(total_min, 2),
                    }
                )

            st.session_state["results"] = results

            st.success(f"Processed {len(results)} items. Videos detected: {len(st.session_state['pending_videos'])}")

            if debug_breakdown and debug_rows:
                with st.expander("Debug: read-time breakdown (Canvas page text + Canvas-hosted docs)", expanded=False):
                    dbg = pd.DataFrame(debug_rows)
                    st.dataframe(dbg, use_container_width=True)

            if debug_do_breakdown and debug_do_rows:
                with st.expander("Debug: DO-time breakdown (complexity + heuristics)", expanded=False):
                    dbg_do = pd.DataFrame(debug_do_rows)
                    st.dataframe(dbg_do, use_container_width=True)

    # 3) Video durations
    st.header("3) Enter video durations (hh:mm:ss)")

    pending = st.session_state.get("pending_videos", {})
    if pending:
        for v_key, meta in list(pending.items()):
            with st.expander(f"{meta['title']} — {meta.get('src','')}"):
                hhmmss = st.text_input(
                    "Duration (hh:mm:ss)",
                    key=f"dur_{v_key}",
                    value=meta.get("hhmmss", "00:00:00"),
                )
                if st.button("💾 Save", key=f"save_{v_key}"):
                    sec = hhmmss_to_seconds(hhmmss)
                    if sec <= 0:
                        st.error("Invalid hh:mm:ss (must be > 00:00:00).")
                    else:
                        meta["hhmmss"] = hhmmss
                        meta["seconds"] = sec
                        st.success("Saved. Totals will update below.")

        # recompute watch minutes per item
        item_seconds: Dict[str, int] = {}
        for meta in pending.values():
            ik = meta.get("item_key")
            if not ik:
                continue
            item_seconds[ik] = item_seconds.get(ik, 0) + int(meta.get("seconds", 0) or 0)

        # apply to results
        for r in st.session_state.get("results", []):
            ik = r.get("item_key")
            sec_total = item_seconds.get(ik, 0)
            watch_min = sec_total / 60.0
            r["watch_min"] = round(watch_min, 2)
            r["total_min"] = round(float(r.get("read_min", 0.0)) + float(r.get("watch_min", 0.0)) + float(r.get("do_min", 0.0)), 2)

    else:
        st.info("No videos detected yet. They’ll appear here after processing items.")

    # 4) Summary
    st.header("4) Workload summary")

    results = st.session_state.get("results", [])
    if not results:
        st.info("No workload results yet. Process items to see estimates.")
        return

    df = pd.DataFrame(results)

    # Ensure module_position exists for older runs
    if "module_position" not in df.columns:
        module_order = {}
        for it in st.session_state.get("items", []):
            mn = it.get("module_name", "")
            pos = it.get("module_position", 0)
            if mn not in module_order or pos < module_order[mn]:
                module_order[mn] = pos
        df["module_position"] = df["module"].map(lambda m: module_order.get(m, 0))

    mod_summary = (
        df.groupby(["module", "module_position"])[["read_min", "watch_min", "do_min", "total_min"]]
        .sum()
        .reset_index()
        .sort_values("module_position")
    )

    grand_totals = {
        "module": "Grand Total",
        "module_position": (mod_summary["module_position"].max() + 1) if len(mod_summary) else 9999,
        "read_min": mod_summary["read_min"].sum(),
        "watch_min": mod_summary["watch_min"].sum(),
        "do_min": mod_summary["do_min"].sum(),
        "total_min": mod_summary["total_min"].sum(),
    }

    mod_summary_with_total = pd.concat([mod_summary, pd.DataFrame([grand_totals])], ignore_index=True)
    mod_summary_display = mod_summary_with_total.drop(columns=["module_position"])

    st.subheader("Per-module totals (minutes)")
    st.dataframe(mod_summary_display, use_container_width=True)

    st.subheader("Item-level details")
    show_cols = ["module", "type", "title", "read_min", "watch_min", "do_min", "total_min", "url"]
    st.dataframe(df[show_cols], use_container_width=True)

    csv = df[show_cols].to_csv(index=False).encode("utf-8")
    st.download_button(
        "Download item-level CSV",
        data=csv,
        file_name="course_load_estimates.csv",
        mime="text/csv",
    )


if __name__ == "__main__":
    main()
